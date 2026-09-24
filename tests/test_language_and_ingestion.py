"""Stages 21 and 22 — multilingual foundations, document ingestion, narration.

These tests do real work on real files. Every document parsed here is generated
in the test itself by the actual library that writes that format, so a passing
DOCX test means python-docx wrote a file and our parser read it back — not that
a fixture we hand-crafted happened to match our own assumptions.

Where a capability is genuinely absent from this environment the test asserts
the *honest failure*: no speech synthesiser means the narration ladder must
degrade to silence and record that it did, not that the assertion is skipped.
"""

from __future__ import annotations

import asyncio
import unittest
from itertools import pairwise
from pathlib import Path
from tempfile import TemporaryDirectory

from vtv.adapters.ingest.documents import (
    MAX_DOCUMENT_BYTES,
    DataParser,
    HtmlParser,
    ParserRegistry,
    TextParser,
)
from vtv.adapters.media import ffmpeg
from vtv.adapters.speech.silent import SilentNarrationProvider, estimate_duration
from vtv.adapters.storage.local import LocalStorageProvider
from vtv.animation import fonts
from vtv.contracts.base import IdPrefix, new_id
from vtv.contracts.errors import Status, ValidationFailed, VTVError
from vtv.contracts.generation import (
    GenerationKind,
    GenerationRequest,
    GenerationResult,
    SpeechParams,
)
from vtv.contracts.language import (
    LANGUAGES,
    Language,
    LanguagePolicy,
    Script,
    detect_script,
    scripts_present,
)
from vtv.contracts.source import BlockKind, InputKind, SourceBlock, TableData
from vtv.contracts.tenancy import SYSTEM_ORGANISATION_ID
from vtv.observability.events import EventSink
from vtv.pipeline.generation import GenerationRouter
from vtv.pipeline.ingestion import (
    IngestionService,
    chart_series_from_table,
    transcript_from_source,
)
from vtv.pipeline.narration import NarrationService, _retime

#: The eighteen languages the platform commits to, from the product brief.
COMMITTED = [
    "te", "hi", "ta", "kn", "ml", "bn", "mr", "gu", "pa", "ur",
    "es", "fr", "de", "pt", "ar", "ja", "ko", "zh",
]


def run(coro):
    return asyncio.run(coro)


def project_id() -> str:
    return new_id(IdPrefix.PROJECT)


# ---------------------------------------------------------------------------
# Language and script
# ---------------------------------------------------------------------------

class LanguageAndScript(unittest.TestCase):
    def test_every_committed_language_is_known(self) -> None:
        """A language we promise must resolve to a real script, not a default.

        This is the test that would have caught shipping Telugu as Latin.
        """
        for code in COMMITTED:
            with self.subTest(code=code):
                self.assertIn(code, LANGUAGES)
                language = Language.parse(code)
                self.assertEqual(language.code, code)
                self.assertIsNot(
                    language.script,
                    Script.UNKNOWN,
                    f"{code} resolved to an unknown script",
                )
                self.assertNotEqual(language.name, "Unknown")

    def test_indian_languages_get_distinct_scripts(self) -> None:
        """Hindi and Telugu are not the same script, and neither is Urdu."""
        self.assertIs(Language.parse("hi").script, Script.DEVANAGARI)
        self.assertIs(Language.parse("te").script, Script.TELUGU)
        self.assertIs(Language.parse("ta").script, Script.TAMIL)
        self.assertIs(Language.parse("kn").script, Script.KANNADA)
        self.assertIs(Language.parse("ml").script, Script.MALAYALAM)
        self.assertIs(Language.parse("bn").script, Script.BENGALI)
        self.assertIs(Language.parse("pa").script, Script.GURMUKHI)
        self.assertIs(Language.parse("gu").script, Script.GUJARATI)
        # Urdu is Arabic script despite being mutually intelligible with Hindi.
        # Picking a font by language rather than script gets exactly this wrong.
        self.assertIs(Language.parse("ur").script, Script.ARABIC)
        self.assertIs(Language.parse("mr").script, Script.DEVANAGARI)

    def test_rtl_is_a_property_of_script(self) -> None:
        self.assertTrue(Language.parse("ar").is_rtl)
        self.assertTrue(Language.parse("ur").is_rtl)
        self.assertTrue(Language.parse("he").is_rtl)
        self.assertFalse(Language.parse("hi").is_rtl)
        self.assertFalse(Language.parse("en").is_rtl)

    def test_cjk_wraps_on_characters(self) -> None:
        """Chinese and Japanese have no spaces; wrapping on them loses the text."""
        self.assertTrue(Script.HAN.wraps_on_characters)
        self.assertTrue(Script.KANA.wraps_on_characters)
        self.assertTrue(Script.THAI.wraps_on_characters)
        self.assertFalse(Script.LATIN.wraps_on_characters)

    def test_complex_scripts_are_flagged(self) -> None:
        """Scripts needing shaping must say so, or they render as broken glyphs."""
        for script in (
            Script.DEVANAGARI,
            Script.TELUGU,
            Script.TAMIL,
            Script.ARABIC,
            Script.BENGALI,
        ):
            with self.subTest(script=script):
                self.assertTrue(script.is_complex)
        self.assertFalse(Script.LATIN.is_complex)

    def test_script_detection_from_real_text(self) -> None:
        cases = {
            "తెలుగు భాష": Script.TELUGU,
            "हिन्दी भाषा": Script.DEVANAGARI,
            "தமிழ் மொழி": Script.TAMIL,
            "ಕನ್ನಡ ಭಾಷೆ": Script.KANNADA,
            "മലയാളം": Script.MALAYALAM,
            "বাংলা ভাষা": Script.BENGALI,
            "ગુજરાતી": Script.GUJARATI,
            "ਪੰਜਾਬੀ": Script.GURMUKHI,
            "العربية": Script.ARABIC,
            "日本語です": Script.KANA,
            "한국어": Script.HANGUL,
            "Hello there": Script.LATIN,
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertIs(detect_script(text), expected)

    def test_mixed_script_text_reports_every_script(self) -> None:
        """Real content mixes scripts: a Telugu sentence with an English brand."""
        found = scripts_present("తెలుగు మరియు English")
        self.assertIn(Script.TELUGU, found)
        self.assertIn(Script.LATIN, found)

    def test_explicit_script_subtag_wins(self) -> None:
        self.assertIs(Language.parse("zh-Hant-TW").script, Script.HAN)
        self.assertEqual(Language.parse("pt-BR").region, "BR")

    def test_unknown_tag_degrades_without_losing_the_original(self) -> None:
        """An unknown tag must not crash the pipeline, and must stay visible."""
        language = Language.parse("xx-YY")
        self.assertEqual(language.tag, "xx-YY")
        self.assertEqual(language.name, "Unknown")

    def test_policy_defaults_every_role_to_the_source(self) -> None:
        policy = LanguagePolicy(source=Language.parse("te"))
        self.assertEqual(policy.understanding_language.code, "te")
        self.assertEqual(policy.caption_language.code, "te")
        self.assertFalse(policy.requires_translation)

    def test_policy_expresses_speak_one_caption_another(self) -> None:
        """Speak Telugu, caption in English — a real request, no special path."""
        policy = LanguagePolicy(
            source=Language.parse("te"), captions=Language.parse("en")
        )
        self.assertTrue(policy.requires_translation)
        self.assertEqual(policy.caption_language.code, "en")
        self.assertIn(Script.TELUGU, policy.all_scripts())
        self.assertIn(Script.LATIN, policy.all_scripts())


class FontCoverage(unittest.TestCase):
    """Honest reporting of what this machine can actually draw."""

    def test_latin_is_production_ready(self) -> None:
        report = fonts.script_support(Script.LATIN)
        self.assertTrue(report["has_font"], "no Latin font found at all")
        self.assertEqual(report["quality"], "good")

    def test_report_separates_ready_from_degraded(self) -> None:
        """The report must never claim a script works when it only half works."""
        report = fonts.support_report()
        ready = set(report["production_ready_scripts"])  # type: ignore[arg-type]
        degraded = set(report["degraded_scripts"])  # type: ignore[arg-type]
        self.assertEqual(ready & degraded, set(), "a script claimed both ways")
        self.assertIn(Script.LATIN.value, ready)

    def test_unsupported_script_yields_an_install_hint(self) -> None:
        """A missing font is an operations problem, so say how to fix it."""
        for script in Script:
            if script is Script.UNKNOWN:
                continue
            support = fonts.script_support(script)
            if not support["has_font"]:
                self.assertTrue(fonts.install_hint(script))

    def test_rtl_is_never_claimed_as_verified(self) -> None:
        """Nobody here can read Arabic. The report must say so, not guess."""
        for script in (Script.ARABIC, Script.HEBREW):
            support = fonts.script_support(script)
            if support["has_font"]:
                self.assertTrue(
                    support.get("note"),
                    "RTL support claimed with no caveat about bidi ordering",
                )


# ---------------------------------------------------------------------------
# Document parsing
# ---------------------------------------------------------------------------

class TextAndMarkdown(unittest.TestCase):
    def test_markdown_structure_survives(self) -> None:
        source = (
            "# The Water Cycle\n\n"
            "Water evaporates from the ocean. It rises and cools.\n\n"
            "## Condensation\n\n"
            "- Vapour becomes droplets\n"
            "- Droplets form clouds\n"
        )
        result = TextParser().parse(
            source.encode(), project_id=project_id(), filename="cycle.md", organisation_id=SYSTEM_ORGANISATION_ID
        )
        kinds = [block.kind for block in result.document.blocks]
        self.assertIn(BlockKind.TITLE, [*kinds, BlockKind.HEADING])
        self.assertIn(BlockKind.LIST_ITEM, kinds)
        outline = result.document.outline()
        self.assertTrue(any("Water Cycle" in title for _, title in outline))
        self.assertTrue(any("Condensation" in title for _, title in outline))

    def test_markdown_table_becomes_data_not_prose(self) -> None:
        source = (
            "# Growth\n\n"
            "| Year | Users |\n|---|---|\n| 2019 | 100 |\n| 2024 | 900 |\n"
        )
        result = TextParser().parse(
            source.encode(), project_id=project_id(), filename="growth.md", organisation_id=SYSTEM_ORGANISATION_ID
        )
        tables = result.document.tables()
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].headers, ["Year", "Users"])
        self.assertTrue(tables[0].is_chartable)

    def test_oversized_input_is_refused_before_parsing(self) -> None:
        with self.assertRaises(VTVError):
            TextParser().parse(
                b"x" * (MAX_DOCUMENT_BYTES + 1),
                project_id=project_id(),
                filename="huge.txt", organisation_id=SYSTEM_ORGANISATION_ID,
            )


class HtmlExtraction(unittest.TestCase):
    def test_navigation_and_scripts_are_dropped(self) -> None:
        """Boilerplate is not content. Narrating a nav bar is a product defect."""
        html = (
            "<html><head><title>Bees</title><style>p{color:red}</style></head>"
            "<body><nav>Home About Contact</nav>"
            "<script>window.track()</script>"
            "<main><h1>Why Bees Matter</h1>"
            "<p>Bees pollinate a third of the food we eat.</p></main>"
            "<footer>Copyright 2024 Example Incorporated</footer></body></html>"
        )
        result = HtmlParser().parse(
            html.encode(), project_id=project_id(), filename="bees.html", organisation_id=SYSTEM_ORGANISATION_ID
        )
        text = result.document.text
        self.assertIn("pollinate", text)
        for boilerplate in ("Home About Contact", "window.track", "Copyright 2024"):
            self.assertNotIn(boilerplate, text)

    def test_title_is_captured(self) -> None:
        result = HtmlParser().parse(
            b"<html><head><title>Tides</title></head><body><p>The moon pulls.</p>"
            b"</body></html>",
            project_id=project_id(),
            filename="tides.html", organisation_id=SYSTEM_ORGANISATION_ID,
        )
        self.assertEqual(result.document.title, "Tides")


class DataFiles(unittest.TestCase):
    def test_csv_becomes_a_chartable_table(self) -> None:
        csv = "Year,Population\n1800,1000000000\n1900,1600000000\n2024,8000000000\n"
        result = DataParser().parse(
            csv.encode(), project_id=project_id(), filename="pop.csv", organisation_id=SYSTEM_ORGANISATION_ID
        )
        tables = result.document.tables()
        self.assertEqual(len(tables), 1)
        series = chart_series_from_table(tables[0])
        self.assertIsNotNone(series)
        assert series is not None
        labels, values, name = series
        # Column zero is the label axis by convention, even when it is numeric:
        # "Year | Population" means years along the bottom.
        self.assertEqual(labels[0], "1800")
        self.assertEqual(values[0], 1_000_000_000.0)
        self.assertEqual(name, "Population")

    def test_json_is_accepted(self) -> None:
        result = DataParser().parse(
            b'[{"city":"Delhi","people":32000000},{"city":"Tokyo","people":37000000}]',
            project_id=project_id(),
            filename="cities.json", organisation_id=SYSTEM_ORGANISATION_ID,
        )
        self.assertIs(result.document.kind, InputKind.JSON)
        self.assertTrue(result.document.tables())


class OfficeFormats(unittest.TestCase):
    """Real files, written by the real libraries, read back by our parsers."""

    def test_docx_round_trip(self) -> None:
        try:
            import docx
        except ImportError:  # pragma: no cover - environment dependent
            self.skipTest("python-docx is not installed")

        with TemporaryDirectory(prefix="vtv-docx-") as scratch:
            path = Path(scratch) / "report.docx"
            document = docx.Document()
            document.add_heading("Quarterly Revenue", level=1)
            document.add_paragraph(
                "Revenue grew steadily through the year as the product matured."
            )
            table = document.add_table(rows=3, cols=2)
            for row, (quarter, value) in enumerate(
                [("Quarter", "Revenue"), ("Q1", "4"), ("Q4", "11")]
            ):
                table.rows[row].cells[0].text = quarter
                table.rows[row].cells[1].text = value
            document.save(path)

            result = ParserRegistry().parse(
                path.read_bytes(), project_id=project_id(), filename="report.docx", organisation_id=SYSTEM_ORGANISATION_ID
            )

        self.assertIs(result.document.kind, InputKind.DOCX)
        self.assertIn("Revenue grew", result.document.text)
        tables = result.document.tables()
        self.assertTrue(tables)
        series = chart_series_from_table(tables[0])
        self.assertIsNotNone(series)
        assert series is not None
        self.assertEqual(series[0], ["Q1", "Q4"])
        self.assertEqual(series[1], [4.0, 11.0])

    def test_pptx_speaker_notes_become_the_narration(self) -> None:
        """The single most valuable judgement in the ingestion path.

        A deck's bullets are on-screen text the author wrote to be *read*. The
        speaker notes are the talk the author wrote to be *said*. Narrating the
        bullets produces a video that reads its own captions aloud.
        """
        try:
            from pptx import Presentation
        except ImportError:  # pragma: no cover - environment dependent
            self.skipTest("python-pptx is not installed")

        with TemporaryDirectory(prefix="vtv-pptx-") as scratch:
            path = Path(scratch) / "deck.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "How Photosynthesis Works"
            slide.placeholders[1].text = "Light\nWater\nCarbon dioxide"
            slide.notes_slide.notes_text_frame.text = (
                "Plants turn sunlight into sugar. First light is absorbed by "
                "chlorophyll, then water is split and carbon dioxide is fixed."
            )
            presentation.save(path)

            result = ParserRegistry().parse(
                path.read_bytes(), project_id=project_id(), filename="deck.pptx", organisation_id=SYSTEM_ORGANISATION_ID
            )

        self.assertIs(result.document.kind, InputKind.PPTX)
        transcript, context = transcript_from_source(result.document)
        spoken = " ".join(segment.text for segment in transcript.segments)
        self.assertIn("chlorophyll", spoken)
        # The bullets were demoted to on-screen material, not narrated.
        self.assertNotIn("Carbon dioxide", spoken)
        self.assertTrue(
            any("Carbon dioxide" in block.text for block in context.apparatus)
        )

    def test_pdf_round_trip(self) -> None:
        try:
            from reportlab.pdfgen import canvas as rlcanvas
        except ImportError:  # pragma: no cover - environment dependent
            self.skipTest("reportlab is not installed")

        with TemporaryDirectory(prefix="vtv-pdf-") as scratch:
            path = Path(scratch) / "paper.pdf"
            pdf = rlcanvas.Canvas(str(path))
            pdf.setFont("Helvetica-Bold", 16)
            pdf.drawString(72, 720, "The Discovery Of Penicillin")
            pdf.setFont("Helvetica", 11)
            pdf.drawString(
                72, 690, "Alexander Fleming discovered penicillin in 1928 in London."
            )
            pdf.drawString(
                72, 674, "It became the first widely used antibiotic in medicine."
            )
            pdf.showPage()
            pdf.save()

            result = ParserRegistry().parse(
                path.read_bytes(), project_id=project_id(), filename="paper.pdf", organisation_id=SYSTEM_ORGANISATION_ID
            )

        self.assertIs(result.document.kind, InputKind.PDF)
        self.assertIn("Fleming", result.document.text)
        # The heading is recognised by shape, so the scene engine sees a topic
        # boundary exactly where the author put one.
        self.assertTrue(
            any(
                block.kind is BlockKind.HEADING and "Penicillin" in block.text
                for block in result.document.blocks
            )
        )
        self.assertTrue(
            all(block.location.page == 1 for block in result.document.blocks)
        )

    def test_registry_routes_by_bytes_not_by_extension(self) -> None:
        """Extensions lie. A PDF named .txt is still a PDF."""
        try:
            from reportlab.pdfgen import canvas as rlcanvas
        except ImportError:  # pragma: no cover
            self.skipTest("reportlab is not installed")

        with TemporaryDirectory(prefix="vtv-sniff-") as scratch:
            path = Path(scratch) / "actually.pdf"
            pdf = rlcanvas.Canvas(str(path))
            pdf.drawString(72, 720, "Sniffing works on magic bytes.")
            pdf.showPage()
            pdf.save()
            data = path.read_bytes()

        result = ParserRegistry().parse(
            data, project_id=project_id(), filename="lying.txt", organisation_id=SYSTEM_ORGANISATION_ID
        )
        self.assertIs(result.document.kind, InputKind.PDF)

    def test_unsupported_bytes_are_refused_clearly(self) -> None:
        with self.assertRaises(VTVError):
            ParserRegistry().for_bytes(b"\x00\x01\x02\x03\xff\xfe", filename="x.bin")


# ---------------------------------------------------------------------------
# Normalisation into a transcript
# ---------------------------------------------------------------------------

class Normalisation(unittest.TestCase):
    def _document(self, blocks: list[SourceBlock]):
        from vtv.contracts.source import SourceDocument

        return SourceDocument(
            organisation_id=SYSTEM_ORGANISATION_ID,
            project_id=project_id(), kind=InputKind.MARKDOWN, blocks=blocks
        )

    def test_headings_get_their_own_segment(self) -> None:
        """An author's section boundary is better evidence than inferred prose."""
        document = self._document(
            [
                SourceBlock(kind=BlockKind.HEADING, text="Part One", level=1),
                SourceBlock(
                    kind=BlockKind.PARAGRAPH,
                    text="This is the body of the first part of the document.",
                ),
            ]
        )
        transcript, _ = transcript_from_source(document)
        self.assertEqual(len(transcript.segments), 2)
        self.assertEqual(transcript.segments[0].speaker, BlockKind.HEADING.value)

    def test_timings_are_labelled_synthetic(self) -> None:
        """Nothing downstream may mistake an estimate for measured audio."""
        document = self._document(
            [SourceBlock(kind=BlockKind.PARAGRAPH, text="Some words to say aloud.")]
        )
        transcript, _ = transcript_from_source(document)
        self.assertEqual(transcript.provider, "synthetic-document")

    def test_segments_never_overlap(self) -> None:
        document = self._document(
            [
                SourceBlock(kind=BlockKind.PARAGRAPH, text=f"Sentence number {n}.")
                for n in range(8)
            ]
        )
        transcript, _ = transcript_from_source(document)
        spans = [segment.span for segment in transcript.segments]
        for earlier, later in pairwise(spans):
            self.assertLessEqual(earlier.end, later.start)

    def test_empty_document_is_refused(self) -> None:
        document = self._document(
            [SourceBlock(kind=BlockKind.FOOTNOTE, text="page 4")]
        )
        with self.assertRaises(ValidationFailed):
            transcript_from_source(document)

    def test_a_table_with_no_numbers_produces_no_chart(self) -> None:
        """A chart of invented numbers is worse than no chart at all."""
        table = TableData(
            headers=["Name", "Role"],
            rows=[["Ada", "Engineer"], ["Grace", "Admiral"]],
        )
        self.assertIsNone(chart_series_from_table(table))
        self.assertFalse(table.is_chartable)

    def test_a_single_row_table_is_not_chartable(self) -> None:
        table = TableData(headers=["Year", "Value"], rows=[["2024", "10"]])
        self.assertFalse(table.is_chartable)

    def test_ingestion_service_emits_an_event(self) -> None:
        events = EventSink()
        seen: list[str] = []
        events.subscribe(lambda event: seen.append(event.name.value))
        service = IngestionService(registry=ParserRegistry(), events=events)
        document, transcript, context = service.ingest(
            b"# Title\n\nA paragraph of real content that can be narrated aloud.\n",
            organisation_id=SYSTEM_ORGANISATION_ID,
            project_id=project_id(),
            filename="note.md",
        )
        self.assertIn("document.ingested", seen)
        self.assertTrue(transcript.segments)
        self.assertTrue(context.outline)
        self.assertEqual(document.parser, "text-1.0")


# ---------------------------------------------------------------------------
# Narration synthesis
# ---------------------------------------------------------------------------

class Narration(unittest.TestCase):
    def setUp(self) -> None:
        self._dir = TemporaryDirectory(prefix="vtv-narration-")
        self.storage = LocalStorageProvider(
            Path(self._dir.name) / "storage", bucket="test"
        )
        self.events = EventSink()

    def tearDown(self) -> None:
        self._dir.cleanup()

    def _transcript(self):
        document_bytes = (
            b"# Rivers\n\n"
            b"A river begins as rainfall on high ground and gathers into streams.\n\n"
            b"Those streams join and carve a valley over thousands of years.\n"
        )
        service = IngestionService(registry=ParserRegistry(), events=self.events)
        _document, transcript, _context = service.ingest(
            document_bytes,
            organisation_id=SYSTEM_ORGANISATION_ID,
            project_id=project_id(),
            filename="rivers.md",
        )
        return transcript

    def test_silence_is_real_audio_of_the_right_length(self) -> None:
        if not ffmpeg.is_available():
            self.skipTest("ffmpeg is not installed")
        provider = SilentNarrationProvider(storage=self.storage)
        request = GenerationRequest(
            organisation_id=SYSTEM_ORGANISATION_ID,
            kind=GenerationKind.SPEECH,
            params=SpeechParams(
                text="One two three four five.",
                segment_texts=["One two three four five."],
            ),
        )
        result = run(provider.synthesize(request))
        self.assertIs(result.status, Status.READY)
        self.assertEqual(len(result.outputs), 1)
        details = result.structured_output or {}
        # has_speech=False is the entire point: the platform must never present
        # a silent track as narration.
        self.assertFalse(details["has_speech"])
        self.assertGreater(float(details["duration_seconds"]), 0.5)
        self.assertEqual(result.cost_usd, 0.0)

    def test_duration_estimate_tracks_word_count(self) -> None:
        short = estimate_duration("Two words", segments=["Two words"])
        long = estimate_duration(
            "word " * 200, segments=[" ".join(["word"] * 200)]
        )
        self.assertLess(short, long)
        self.assertGreater(long, 60.0)

    def test_no_synthesiser_degrades_to_silence_and_says_so(self) -> None:
        """The honest failure. This environment has no TTS; assert what happens."""
        if not ffmpeg.is_available():
            self.skipTest("ffmpeg is not installed")
        service = NarrationService(
            router=GenerationRouter(events=self.events),
            events=self.events,
            silent_fallback=SilentNarrationProvider(storage=self.storage),
        )
        spoken = run(service.synthesise(self._transcript()))
        self.assertFalse(spoken.has_speech)
        self.assertEqual(len(spoken.degradations), 1)
        self.assertEqual(spoken.degradations[0].to_strategy, "silent_narration")
        self.assertEqual(spoken.provider, "silent-narration")

    def test_no_synthesiser_and_no_fallback_refuses(self) -> None:
        """A deployment may prefer failure to a mute video. It must be able to."""
        service = NarrationService(
            router=GenerationRouter(events=self.events),
            events=self.events,
            silent_fallback=None,
        )
        with self.assertRaises(VTVError):
            run(service.synthesise(self._transcript()))

    def test_transcript_is_rescaled_onto_the_audio_that_exists(self) -> None:
        """Captions must match what the viewer hears, not what we guessed."""
        transcript = self._transcript()
        original = transcript.span
        assert original is not None
        retimed = _retime(transcript, original.end * 2.0)
        new_span = retimed.span
        assert new_span is not None
        self.assertAlmostEqual(new_span.end, original.end * 2.0, delta=0.2)
        self.assertEqual(len(retimed.segments), len(transcript.segments))
        # Scaled word timings would be fiction at word granularity.
        self.assertTrue(all(not s.words for s in retimed.segments))

    def test_retiming_is_skipped_when_the_estimate_was_right(self) -> None:
        transcript = self._transcript()
        span = transcript.span
        assert span is not None
        self.assertIs(_retime(transcript, span.end), transcript)

    def test_a_provider_that_lies_about_duration_is_rejected(self) -> None:
        """A wrong duration desynchronises every caption in the video."""
        storage = self.storage

        class LyingProvider:
            async def synthesize(self, request: GenerationRequest) -> GenerationResult:
                ref = await storage.put(
                    key="lying/narration.wav",
                    data=b"RIFF",
                    content_type="audio/wav",
                )
                return GenerationResult(
                    request_id=request.request_id,
                    cache_key=request.cache_key(),
                    status=Status.READY,
                    provider="liar",
                    outputs=[ref],
                    structured_output={"duration_seconds": 0.0},
                )

        service = NarrationService(
            router=GenerationRouter(events=self.events),
            events=self.events,
            silent_fallback=LyingProvider(),
        )
        with self.assertRaises(VTVError):
            run(service.synthesise(self._transcript()))


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
